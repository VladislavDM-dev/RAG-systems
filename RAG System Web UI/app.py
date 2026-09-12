"""Web-интерфейс для RAG System.

Отдельный Flask-сервис: страницы чата/админки/теста парсера,
проксирование /rag/* на HTTP API сервиса RAG System.
Чат работает через потоковый API /rag/agent/stream, история и избранное
хранятся на сервере по пользователям (см. DATA_DIR, users.json).
"""
import json
import logging
import os
import re
import difflib
import threading

import requests
from flask import Flask, Response, jsonify, redirect, render_template, request, session, send_file, url_for

from docx import Document

import memo_docx

app = Flask(__name__)
app.json.ensure_ascii = False
app.secret_key = os.environ.get("SECRET_KEY", "rag-web-ui-dev-secret")
# Версия JS-мастера — меняется при правках, чтобы браузер не держал кеш.
app.config["MEMO_JS_VERSION"] = "20260910-31"

log = logging.getLogger("werkzeug")
log.setLevel(logging.ERROR)

# Адрес API сервиса RAG System (порт 11455).
RAG_API_URL = os.environ.get("RAG_API_URL", "http://127.0.0.1:11455/rag")
PORT = int(os.environ.get("PORT", 80))

# Пользователи и история на сервере.
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
USERS_FILE = os.environ.get("USERS_FILE", os.path.join(BASE_DIR, "users.json"))
DATA_DIR = os.environ.get("DATA_DIR", os.path.join(BASE_DIR, "data"))
DEFAULT_ADMIN_PASSWORD = os.environ.get("DEFAULT_ADMIN_PASSWORD", "admin123")

AGENT_SYSTEM_PROMPT = (
    "Ты — RAG-ассистент, отвечающий на основе предоставленных документов. "
    "Отвечай на вопрос, опираясь только на информацию из контекста ниже. "
    "Дай РАЗВЁРНУТЫЙ и подробный ответ: раскрой тему из документа максимально полно, "
    "приведи все релевантные факты, цифры, условия и детали, которые есть в контексте, "
    "структурируй ответ абзацами и, где уместно, списками. "
    "Не выдумывай факты и не добавляй информацию, которой нет в контексте. "
    "Если в контексте нет нужных данных, честно скажи, что информации нет."
)

MEMO_SYSTEM_PROMPT = (
    "Ты — RAG-ассистент, который составляет официальную служебную записку "
    "на основе предоставленных документов. "
    "Оформи текст как СЛУЖЕБНУЮ ЗАПИСКУ: в начале — адресат (Кому), составитель (От), "
    "заголовок с темой, затем основная часть с развёрнутым изложением предмета записки "
    "(факты, цифры, даты, реквизиты документов из контекста), в конце — вывод и предложение. "
    "Стиль официально-деловой, на русском языке, с абзацами. "
    "Опирайся ТОЛЬКО на информацию из контекста ниже; "
    "не выдумывай факты и не добавляй детали, которых нет в контексте. "
    "Если нужных данных нет, честно укажи, какой информации недостаёт."
)

COMMAND_CREATE = "/create"

_lock = threading.Lock()

# Значения конфигурации по умолчанию (fallback, если RAG System недоступен).
DEFAULT_CONFIG = {
    "version": "—",
    "data_bd": "—",
    "document_count": 0,
    "provider": "ollama",
    "llm_url": "http://127.0.0.1:11434",
    "api_key": "",
    "llm": "qwen3:4b",
    "llm_embed": "nomic-embed-text-v2-moe:latest",
    "llm_keywords": "llama3.2:3b",
    "context_window": 4096,
    "chunk_size": 256,
    "chunk_overlap": 50,
    "similarity_top_k": 10,
}


# ---------------------------------------------------------------------------
#  Пользователи и авторизация
# ---------------------------------------------------------------------------
def _ensure_users_file():
    """Создаёт users.json с учёткой admin, если файла ещё нет."""
    if os.path.exists(USERS_FILE):
        return
    os.makedirs(os.path.dirname(USERS_FILE), exist_ok=True)
    with _lock:
        _write_users({
            "users": {
                "admin": DEFAULT_ADMIN_PASSWORD,
            }
        })
    print(f"[auth] Создан users.json. Логин: admin, пароль: {DEFAULT_ADMIN_PASSWORD}")


def _read_users() -> dict:
    try:
        with open(USERS_FILE, "r", encoding="utf-8-sig") as f:
            data = json.load(f) or {"users": {}}
    except (OSError, ValueError):
        data = {"users": {}}
    # Совместимость с хэшированными паролями (миграция не нужна — проверка ниже).
    return data


def _write_users(data: dict):
    with open(USERS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def _hash_like(value: str) -> bool:
    """Пароль хранится как werkzeug-хэш (scrypt:/pbkdf2:...) или открытым текстом."""
    return value.startswith(("scrypt:", "pbkdf2:"))


def authenticate(username: str, password: str) -> bool:
    if not username:
        return False
    users = _read_users().get("users", {})
    stored = users.get(username)
    if stored is None:
        return False
    try:
        from werkzeug.security import check_password_hash
    except ImportError:
        check_password_hash = None
    if _hash_like(stored) and check_password_hash is not None:
        return check_password_hash(stored, password)
    return stored == password


def register_user(username: str, password: str) -> str | None:
    """Регистрирует нового пользователя. Возвращает текст ошибки или None."""
    username = (username or "").strip()
    if not re.fullmatch(r"[A-Za-z0-9_.-]{3,32}", username):
        return "Логин: 3-32 символа, только буквы, цифры, точка, дефис, подчёркивание"
    if len(password) < 4:
        return "Пароль должен быть не короче 4 символов"
    users = _read_users()
    users.setdefault("users", {})
    if username in users["users"]:
        return "Такой логин уже существует"
    with _lock:
        users["users"][username] = password
        _write_users(users)
    return None


def list_users() -> list[dict]:
    users = _read_users().get("users", {})
    return [
        {"username": name, "password": pwd}
        for name, pwd in sorted(users.items())
    ]


def delete_user(username: str) -> bool:
    if username == "admin":
        return False
    users = _read_users()
    if username in users.get("users", {}):
        with _lock:
            del users["users"][username]
            _write_users(users)
        return True
    return False


def set_user_password(username: str, password: str) -> bool:
    if password is None or len(password) < 4:
        return False
    users = _read_users()
    if username not in users.get("users", {}):
        return False
    with _lock:
        users["users"][username] = password
        _write_users(users)
    return True


def require_login(view):
    """Декоратор: редиректит на /login, если пользователь не авторизован."""
    def wrapper(*args, **kwargs):
        if not session.get("user"):
            return redirect(url_for("login_page"))
        return view(*args, **kwargs)
    wrapper.__name__ = view.__name__
    return wrapper


def require_admin(view):
    """Декоратор: только для пользователя 'admin'."""
    def wrapper(*args, **kwargs):
        if session.get("user") != "admin":
            return jsonify({"error": "Доступно только администратору"}), 403
        return view(*args, **kwargs)
    wrapper.__name__ = view.__name__
    return wrapper


@app.route("/login", methods=["GET"])
def login_page():
    if session.get("user"):
        return redirect(url_for("index"))
    next_url = request.args.get("next", "/")
    return render_template("login.html", next_url=next_url)


@app.route("/login", methods=["POST"])
def login_submit():
    username = (request.form.get("username") or "").strip()
    password = request.form.get("password") or ""
    if authenticate(username, password):
        session["user"] = username
        next_url = request.form.get("next") or "/"
        safe = next_url if next_url.startswith("/") else "/"
        return redirect(safe)
    return render_template("login.html", error="Неверный логин или пароль", next_url=request.form.get("next") or "/")


@app.route("/register", methods=["GET"])
def register_page():
    if session.get("user"):
        return redirect(url_for("index"))
    return render_template("register.html")


@app.route("/register", methods=["POST"])
def register_submit():
    username = request.form.get("username") or ""
    password = request.form.get("password") or ""
    password2 = request.form.get("password2") or ""
    if password != password2:
        return render_template("register.html", error="Пароли не совпадают",
                               username=username)
    error = register_user(username, password)
    if error:
        return render_template("register.html", error=error, username=username)
    session["user"] = username
    return redirect(url_for("index"))


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login_page"))


# ---------------------------------------------------------------------------
#  История пользователя на сервере
# ---------------------------------------------------------------------------
def _user_file(username: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", username) or "user"
    return os.path.join(DATA_DIR, f"{safe}.json")


def _read_history(username: str) -> dict:
    path = _user_file(username)
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return {
            "conversations": data.get("conversations", []),
            "answers": data.get("answers", []),
        }
    except (OSError, ValueError):
        return {"conversations": [], "answers": []}


def _write_history(username: str, data: dict):
    os.makedirs(DATA_DIR, exist_ok=True)
    path = _user_file(username)
    with _lock:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)


def _upsert(collection, item, key="id"):
    for i, existing in enumerate(collection):
        if existing.get(key) == item.get(key):
            collection[i] = item
            return
    collection.insert(0, item)


def _remove(collection, item_id, key="id"):
    idx = next((i for i, x in enumerate(collection) if x.get(key) == item_id), None)
    if idx is not None:
        collection.pop(idx)


@app.route("/api/history", methods=["GET"])
@require_login
def api_history():
    user = session["user"]
    return jsonify(_read_history(user))


@app.route("/api/conversations", methods=["PUT"])
@require_login
def api_upsert_conversation():
    user = session["user"]
    item = request.get_json(silent=True) or {}
    if not item.get("id"):
        return jsonify({"error": "Поле 'id' обязательно"}), 400
    data = _read_history(user)
    _upsert(data["conversations"], item)
    _write_history(user, data)
    return jsonify({"ok": True})


@app.route("/api/conversations/<cid>", methods=["DELETE"])
@require_login
def api_delete_conversation(cid):
    user = session["user"]
    data = _read_history(user)
    _remove(data["conversations"], cid)
    _write_history(user, data)
    return jsonify({"ok": True})


@app.route("/api/answers", methods=["PUT"])
@require_login
def api_upsert_answer():
    user = session["user"]
    item = request.get_json(silent=True) or {}
    if not item.get("id"):
        return jsonify({"error": "Поле 'id' обязательно"}), 400
    data = _read_history(user)
    _upsert(data["answers"], item)
    _write_history(user, data)
    return jsonify({"ok": True})


@app.route("/api/answers/<aid>", methods=["DELETE"])
@require_login
def api_delete_answer(aid):
    user = session["user"]
    data = _read_history(user)
    _remove(data["answers"], aid)
    _write_history(user, data)
    return jsonify({"ok": True})


@app.route("/api/history/clear", methods=["POST"])
@require_login
def api_clear_history():
    user = session["user"]
    _write_history(user, {"conversations": [], "answers": []})
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
#  Страницы
# ---------------------------------------------------------------------------
@app.route("/")
@require_login
def index():
    """Чат-интерфейс. Проксирует запросы на /rag/agent/stream."""
    return render_template("index.html", user=session["user"])


@app.route("/admin")
@require_admin
def admin():
    """Админ-панель: инфо, настройки LLM, папки, пользователи, синхронизация."""
    status = get_json("engine/status")
    cfg = get_json("config")
    models = get_json("models").get("models") or {}

    def fallback(key, default):
        return cfg.get(key, default)

    context = dict(DEFAULT_CONFIG)
    context.update({
        "version": status.get("version") or DEFAULT_CONFIG["version"],
        "data_bd": status.get("index_updated") or DEFAULT_CONFIG["data_bd"],
        "document_count": status.get("document_count") or DEFAULT_CONFIG["document_count"],
        "provider": fallback("provider", DEFAULT_CONFIG["provider"]),
        "api_key": fallback("api_key", DEFAULT_CONFIG["api_key"]),
        "llm_url": fallback("url", DEFAULT_CONFIG["llm_url"]),
        "llm": fallback("llm", DEFAULT_CONFIG["llm"]),
        "llm_embed": fallback("llm_embed", DEFAULT_CONFIG["llm_embed"]),
        "llm_keywords": fallback("llm_keywords", DEFAULT_CONFIG["llm_keywords"]),
        "context_window": fallback("context_window", DEFAULT_CONFIG["context_window"]),
        "chunk_size": fallback("chunk_size", DEFAULT_CONFIG["chunk_size"]),
        "chunk_overlap": fallback("chunk_overlap", DEFAULT_CONFIG["chunk_overlap"]),
        "similarity_top_k": fallback("similarity_top_k", DEFAULT_CONFIG["similarity_top_k"]),
        "llm_thinking": models.get("thinking", []),
        "llm_embedding": models.get("embedding", []),
        "llm_all": models.get("all", []),
        "users": list_users(),
        "memo_path": memo_docx.TEMPLATE_PATH,
        "memo_updated": memo_docx.get_template_info()["updated"],
        "memo_fields": ", ".join(memo_docx.get_template_fields()),
    })
    return render_template("admin.html", **context)


# Админ: управление пользователями.
@app.route("/api/admin/users", methods=["GET"])
@require_admin
def api_users_list():
    return jsonify({"users": list_users()})


@app.route("/api/admin/users", methods=["POST"])
@require_admin
def api_users_add():
    data = request.get_json(silent=True) or {}
    username = data.get("username") or ""
    password = data.get("password") or ""
    error = register_user(username, password)
    if error:
        return jsonify({"error": error}), 400
    return jsonify({"ok": True, "users": list_users()})


@app.route("/api/admin/users/<username>", methods=["DELETE"])
@require_admin
def api_users_delete(username):
    ok = delete_user(username)
    if not ok:
        return jsonify({"error": "Не удалось удалить пользователя (admin нельзя)"}), 400
    return jsonify({"ok": True, "users": list_users()})


@app.route("/api/admin/users/<username>/password", methods=["PUT"])
@require_admin
def api_users_set_password(username):
    data = request.get_json(silent=True) or {}
    password = data.get("password") or ""
    ok = set_user_password(username, password)
    if not ok:
        return jsonify({"error": "Не удалось изменить пароль"}), 400
    return jsonify({"ok": True, "users": list_users()})


# Служебная записка: шаблон и выгрузка .docx.
@app.route("/api/memo/template", methods=["GET"])
def api_memo_template():
    """Возвращает описание полей шаблона служебной записки.

    Структура определяется сервером из самого файла шаблона (docx_templates/mmo.dotx).
    """
    info = memo_docx.get_template_info()
    return jsonify({
        "fields": [
            {
                "key": key,
                "label": memo_docx.FIELD_LABELS.get(key, key.replace("_", " ").title()),
                "hint": memo_docx.FIELD_HINTS.get(key, ""),
            }
            for key in info["fields"]
        ],
        "path": info["path"],
        "updated": info["updated"],
    })


@app.route("/api/admin/memo/template/reload", methods=["POST"])
@require_admin
def api_admin_memo_template_reload():
    """Перечитывает поля шаблона с диска (кнопка админа «Обновить шаблоны»)."""
    fields = memo_docx.reload_template_fields()
    return jsonify({
        "ok": True,
        "fields": fields,
        "path": memo_docx.TEMPLATE_PATH,
        "updated": memo_docx.get_template_info()["updated"],
    })


@app.route("/api/admin/memo/template", methods=["POST"])
@require_admin
def api_admin_memo_template_upload():
    """Загружает новый файл шаблона служебной записки (.dotx)."""
    if "file" not in request.files:
        return jsonify({"error": "Файл не передан (поле 'file')"}), 400
    file = request.files["file"]
    if not file.filename or not file.filename.lower().endswith(".dotx"):
        return jsonify({"error": "Нужен файл шаблона .dotx"}), 400
    fields = memo_docx.save_template_file(file.read())
    return jsonify({
        "ok": True,
        "fields": fields,
        "path": memo_docx.TEMPLATE_PATH,
        "updated": memo_docx.get_template_info()["updated"],
    })


@app.route("/api/memo/docx", methods=["POST"])
@require_login
def api_memo_docx():
    """Собирает .docx служебной записки по шаблону и отдаёт файлом."""
    data = request.get_json(silent=True) or {}
    try:
        buf = memo_docx.generate_memo_docx(data)
    except FileNotFoundError as e:
        return jsonify({"error": str(e)}), 500
    return send_file(
        buf,
        as_attachment=True,
        download_name="sluzhebnaya_zapiska.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@app.route("/test-parser")
def test_parser():
    """Тест парсера файлов."""
    return render_template("test_parser.html")


# Список служебных записок для простого запроса (chron: до RAG не дотягивает).
_MEMO_FOLDER = os.path.abspath(os.path.join(BASE_DIR, "..", "out"))
_MEMOS_CACHE = {"ts": 0.0, "memos": None}
_MEMOS_CACHE_TTL = 10
_MEMO_STOP = {
    "покажи", "показывай", "найди", "нужно", "дай", "выдай", "дать",
    "список", "все", "всё", "полный", "служебные", "служебную",
    "записки", "записку", "записка", "документ", "документы", "какие",
    "есть", "мне", "их", "содержимое", "содержимым", "по", "про",
    "о", "и", "в", "на", "с", "не", "то", "что", "где", "как", "это",
    "очень", "просто", "те", "которых", "который", "которые", "у",
}


def extract_memo_body(text: str) -> str:
    """Вырезает тело служебной записки: без шапки (Кому/№/дата) и подписи.

    Формат готового документа: «Кому: …», «Служебная записка № … от …»,
    затем тело, затем «С уважением, …». Возвращает текст между ними,
    при невозможности разобрать — весь текст.
    """
    lines = (text or "").splitlines()

    def line_key(idx: int) -> str:
        return lines[idx].strip().lower()

    start = next((i for i, ln in enumerate(lines)
                  if line_key(i).startswith("служебная записка №")), None)
    end = next((i for i, ln in enumerate(lines)
                if line_key(i).startswith("с уважением")), None)

    if start is None:
        return (text or "").strip()
    body_lines = lines[start + 1: end]
    while body_lines and not body_lines[0].strip():
        body_lines.pop(0)
    while body_lines and not body_lines[-1].strip():
        body_lines.pop()
    body = "\n".join(body_lines).strip()
    return body or (text or "").strip()


def _extract_contact(header_parts: list) -> tuple:
    """Из текста колонтитула достаёт (ФИО, телефон): телефон — по маске,
    ФИО — первая непустая строка без признаков телефона."""
    phone = next((p.strip() for p in header_parts
                  if re.search(r"\+?7\s*[\d\(\)\-]{5,}", p)), "")
    executor = next((p.strip() for p in header_parts
                     if p.strip() and not re.search(r"\+?7\s*[\d\(\)\-]{5,}", p)), "")
    return executor or "", phone or ""


def _all_memos_text() -> list:
    """Все служебные записки (docx из каталога ../out) с содержимым."""
    import time as _t
    now = _t.time()
    cached = _MEMOS_CACHE["memos"]
    if cached is not None and (now - _MEMOS_CACHE["ts"]) < _MEMOS_CACHE_TTL:
        return cached

    memos = []
    if os.path.isdir(_MEMO_FOLDER):
        for name in sorted(os.listdir(_MEMO_FOLDER)):
            if not name.lower().endswith(".docx") or name.startswith("~"):
                continue
            full = os.path.join(_MEMO_FOLDER, name)
            try:
                doc = Document(full)
                parts = [p.text or "" for p in doc.paragraphs]
                for tbl in doc.tables:
                    for row in tbl.rows:
                        row_text = " | ".join(c.text or "" for c in row.cells)
                        if row_text.strip():
                            parts.append(row_text)
                text = "\n".join(parts).strip()

                # Исполнитель и телефон живут в колонтитуле (footer) документа.
                header_parts = []
                for sec in doc.sections:
                    for part in (sec.footer, sec.header,
                                 sec.first_page_footer, sec.first_page_header):
                        try:
                            for p in part.paragraphs:
                                t = (p.text or "").strip()
                                if t:
                                    header_parts.append(t)
                        except Exception:
                            continue
                executor, phone = _extract_contact(header_parts)

                if not text:
                    continue
            except Exception as e:
                print(f"[WARN] Не удалось прочитать служебку {full}: {e}")
                continue
            memos.append({
                "file_name": name,
                "file_path": full,
                "text": text,
                "body": extract_memo_body(text),
                "executor": executor,
                "phone": phone,
            })
    _MEMOS_CACHE["memos"] = memos
    _MEMOS_CACHE["ts"] = now
    return memos


def _filter_memos(memos: list, query: str) -> list:
    """Только записки, где есть значимые слова запроса (с опечатками)."""
    query = (query or "").strip().lower()
    if not query:
        return memos

    words = [w for w in query.split() if len(w) > 1 and w not in _MEMO_STOP]
    if not words:
        return memos

    doc_words = []
    for m in memos:
        toks = set(re.findall(r"[а-яёa-z]+", m["text"].lower()))
        toks = {w for w in toks if len(w) > 2 and w not in _MEMO_STOP}
        doc_words.append(toks)

    def _matches(word, text_words):
        if word in text_words:
            return True
        for tw in text_words:
            if abs(len(tw) - len(word)) > 3:
                continue
            if difflib.SequenceMatcher(None, word, tw).ratio() >= 0.8:
                return True
            common = 0
            for a, b in zip(word, tw):
                if a == b:
                    common += 1
                else:
                    break
            if common >= 6:
                return True
        return False

    scored = []
    for m, tw in zip(memos, doc_words):
        hits = sum(1 for w in words if _matches(w, tw))
        if hits:
            scored.append((hits, m))
    scored.sort(key=lambda pair: (-pair[0], pair[1]["file_name"]))
    return [m for _, m in scored]


@app.route("/api/memos", methods=["GET"])
def api_memos():
    """Служебные записки с содержимым; q — фильтр по тексту запроса."""
    query = request.args.get("q", "")
    memos = _filter_memos(_all_memos_text(), query)
    return jsonify({"memos": memos})


def get_json(endpoint, timeout=5):
    """GET к API RAG System. При недоступности возвращает пустой dict."""
    try:
        resp = requests.get(f"{RAG_API_URL}/{endpoint}", timeout=timeout)
        if resp.status_code == 200:
            return resp.json()
    except requests.RequestException:
        pass
    return {}


def _dedupe_sources(sources):
    """Убирает дубли документов в списке источников, оставляя лучший score."""
    seen = {}
    for src in sources or []:
        key = src.get("file_path") or src.get("file_name") or ""
        if not key:
            continue
        if key not in seen or float(src.get("score") or 0) > float(seen[key].get("score") or 0):
            seen[key] = src
    return list(seen.values())


def generate_stream(resp):
    """Передаёт тело streaming-ответа (SSE) из RAG System в браузер.

    В событиях sources/done дубли документов (одинаковый file_path) удаляются.
    """
    buffer = b""
    try:
        for chunk in resp.iter_content(chunk_size=4096):
            buffer += chunk
            while b"\n\n" in buffer:
                line, buffer = buffer.split(b"\n\n", 1)
                if b'"type": "sources"' in line or b'"type": "done"' in line:
                    yield _dedupe_stream_line(line) + b"\n\n"
                else:
                    yield line + b"\n\n"
        if buffer.strip():
            if b'"type": "sources"' in buffer or b'"type": "done"' in buffer:
                yield _dedupe_stream_line(buffer)
            else:
                yield buffer
    finally:
        resp.close()


def _dedupe_stream_line(line):
    """Дедуплицирует источники в одной SSE-строке (data: {...})."""
    prefix, sep, payload = line.partition(b"data: ")
    if not sep:
        return line
    try:
        data = json.loads(payload.decode("utf-8"))
    except (ValueError, UnicodeDecodeError):
        return line
    if data.get("type") == "sources":
        data["data"] = _dedupe_sources(data.get("data"))
    elif data.get("type") == "done":
        data["sources"] = _dedupe_sources(data.get("sources"))
    return prefix + sep + json.dumps(data, ensure_ascii=False).encode("utf-8")


@app.route("/rag/<path:path>", methods=["GET", "POST", "PUT", "PATCH", "DELETE"])
@require_login
def proxy(path):
    """Проксирование запросов на HTTP API RAG System.

    Пути, оканчивающиеся на /stream, передаются как SSE (stream=True).
    """
    upstream = f"{RAG_API_URL}/{path}"
    headers = {"Content-Type": request.content_type or "application/json"}
    params = request.args.to_dict(flat=True)
    body = request.get_data()

    # Для /rag/agent/stream подставляем системный промпт, если клиент его не задал.
    # Команда /create <тема> переключает на промпт генерации служебной записки.
    if path == "agent/stream" and request.method == "POST":
        try:
            payload = json.loads(body.decode("utf-8")) if body else {}
            if isinstance(payload, dict):
                prompt = (payload.get("prompt") or "").strip()
                memo = False
                if prompt.startswith(COMMAND_CREATE):
                    topic = prompt[len(COMMAND_CREATE):].strip()
                    if topic:
                        payload["prompt"] = (
                            f"Составь официальную служебную записку по теме: {topic}. "
                            "Используй данные из контекста документов."
                        )
                    memo = True
                if not payload.get("system_prompt"):
                    payload["system_prompt"] = (
                        MEMO_SYSTEM_PROMPT if memo else AGENT_SYSTEM_PROMPT
                    )
                    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
                elif memo:
                    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        except (ValueError, UnicodeDecodeError):
            pass

    try:
        if path.endswith("/stream"):
            resp = requests.request(
                request.method, upstream, params=params, data=body,
                headers=headers, stream=True, timeout=(5, None),
            )
            if resp.status_code >= 400:
                return Response(resp.content, status=resp.status_code,
                                content_type="application/json")
            return Response(generate_stream(resp), mimetype="text/event-stream")

        resp = requests.request(
            request.method, upstream, params=params, data=body,
            headers=headers, timeout=(5, 120),
        )
    except requests.RequestException as e:
        return jsonify({"error": "RAG System API недоступен", "detail": str(e)}), 502

    content_type = resp.headers.get("Content-Type", "application/json")
    return Response(resp.content, status=resp.status_code, content_type=content_type)


if __name__ == "__main__":
    _ensure_users_file()
    print(f"RAG System Web UI: http://127.0.0.1:{PORT}  ->  API {RAG_API_URL}")
    print("   Порт 80 может требовать прав администратора.")
    print(f"   Пользователи: {USERS_FILE}, история: {DATA_DIR}")
    app.run(debug=False, host="0.0.0.0", port=PORT, use_reloader=False)